import os
import runpy
import shutil
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfWriter

import check_invoices
import extractor
from check_invoices import GREEN, RED, Extraction, parse_amount, parse_date, run
from make_test_invoices import line_items

ROOT = Path(__file__).parent.parent
RUN_DATE = date(2026, 10, 15)
KEY = pd.read_csv(ROOT / "answer_key.csv", dtype=str, keep_default_na=False).set_index("file")
FIELDS = [
    "document_type", "supplier_name", "gst_number", "invoice_number", "invoice_date", "due_date",
    "buyer_name", "buyer_identifier", "reference", "work_order", "site", "subtotal", "gst", "total",
    "bank_account",
]


def fake_extract(pdf):
    """What a perfect model would return: the answer key values as source text."""
    row = KEY.loc[Path(pdf).name]

    def field(text):
        return {"source_text": text or None, "confidence": "high"}

    out = {f: field(row[f]) for f in FIELDS}
    out["line_items"] = [{"description": field(d), "amount": field(a)} for d, a in line_items(row)]
    return out


def check(tmp_path, files, extract=fake_extract):
    folder = tmp_path / "invoices"
    folder.mkdir()
    for f in files:
        shutil.copy(ROOT / "invoices" / f, folder)
    out = tmp_path / "review.xlsx"
    run(folder, out, extract, RUN_DATE)
    return read_review(out)


def read_review(path):
    header, *rows = load_workbook(path).active.iter_rows(values_only=True)
    return [dict(zip(header, r)) for r in rows]


def test_clean_invoice_gives_one_looks_fine_row_in_xero_order(tmp_path):
    [row] = check(tmp_path, ["01-clean.pdf"])
    assert list(row) == [
        "File", "Status", "Flags", "Notes", "Suggested property", "Property reason",
        "Supplier", "Date", "Due date", "Bill reference", "Line items", "Site", "Subtotal", "GST", "Total",
    ]
    assert row["Status"] == "Looks fine"
    assert row["Flags"] is None
    assert row["Supplier"] == "Southbridge Electrical Limited"
    assert row["Date"] == datetime(2026, 10, 6)
    assert row["Due date"] == datetime(2026, 11, 20)
    assert row["Bill reference"] == "INV-2147"
    assert row["Line items"] == KEY.loc["01-clean.pdf", "line_items"]
    assert row["Site"] == "Kowhai House"
    amounts = [Decimal(str(row[c])) for c in ("Subtotal", "GST", "Total")]
    assert amounts == [Decimal("1026.00"), Decimal("153.90"), Decimal("1179.90")]


def test_source_text_not_in_pdf_leaves_field_blank_and_flags_it(tmp_path):
    def extract(pdf):
        out = fake_extract(pdf)
        out["invoice_number"]["source_text"] = "INV-9999"
        return out

    [row] = check(tmp_path, ["01-clean.pdf"], extract)
    assert row["Status"] == "Needs you"
    assert row["Bill reference"] is None
    assert row["Flags"] == "Source text not found: invoice number 'INV-9999' is not in the PDF text."


def test_pdf_with_no_text_gives_one_unreadable_flag_without_calling_extract(tmp_path):
    folder = tmp_path / "invoices"
    folder.mkdir()
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.write(folder / "scan.pdf")

    def extract(pdf):
        raise AssertionError("extract must not be called")

    run(folder, tmp_path / "review.xlsx", extract, RUN_DATE)
    [row] = read_review(tmp_path / "review.xlsx")
    assert row["Status"] == "Needs you"
    assert row["Flags"].startswith("Unreadable:")
    assert "\n" not in row["Flags"]


def test_extract_raising_gives_one_could_not_read_flag_with_the_error(tmp_path):
    def extract(pdf):
        raise RuntimeError("Connection timed out")

    [row] = check(tmp_path, ["01-clean.pdf"], extract)
    assert row["Status"] == "Needs you"
    assert row["Flags"].startswith("Could not read:")
    assert "Connection timed out" in row["Flags"]
    assert "\n" not in row["Flags"]
    assert row["Supplier"] is None


def test_reader_error_never_shows_the_api_key(tmp_path):
    def extract(pdf):
        raise RuntimeError("Incorrect API key provided: sk-proj-****abcd.")

    [row] = check(tmp_path, ["01-clean.pdf"], extract)
    assert "sk-" not in row["Flags"]
    assert "[key hidden]" in row["Flags"]


def run_command_line(tmp_path, monkeypatch, *options):
    """python check_invoices.py <folder with #1> [options], with OpenAI and Excel faked."""
    calls, opened = [], []

    class FakeOpenAI:
        def __init__(self):
            self.responses = self

        def parse(self, **request):
            calls.append(request)
            parsed = Extraction.model_validate(fake_extract(ROOT / "invoices" / "01-clean.pdf"))
            return SimpleNamespace(output_parsed=parsed)

    monkeypatch.setattr(extractor, "OpenAI", FakeOpenAI)
    monkeypatch.setattr(extractor, "load_dotenv", lambda: None)
    monkeypatch.setenv("OPENAI_MODEL", "model-from-setting")
    folder = tmp_path / "invoices"
    folder.mkdir()
    shutil.copy(ROOT / "invoices" / "01-clean.pdf", folder)
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(os, "startfile", opened.append, raising=False)
    monkeypatch.setattr(sys, "argv", ["check_invoices.py", str(folder), *options])
    runpy.run_path(str(ROOT / "check_invoices.py"), run_name="__main__")
    return calls, opened


def test_command_line_on_invoice_1_with_a_fake_openai_is_looks_fine(tmp_path, monkeypatch):
    calls, opened = run_command_line(tmp_path, monkeypatch)

    [row] = read_review(tmp_path / "review.xlsx")
    assert row["Status"] == "Looks fine", row["Flags"]
    [request] = calls
    assert request["model"] == "model-from-setting"
    assert request["text_format"] is Extraction
    assert "INV-2147" in request["input"]
    assert opened == []


def test_open_option_opens_review_xlsx_when_the_run_finishes(tmp_path, monkeypatch):
    _, opened = run_command_line(tmp_path, monkeypatch, "--open")
    assert [Path(p) for p in opened] == [tmp_path / "review.xlsx"]


def test_review_is_easy_to_read_when_it_opens(tmp_path):
    def extract(pdf):
        out = fake_extract(pdf)
        if pdf.name == "02-wrong-gst.pdf":
            out["invoice_number"]["source_text"] = "INV-9999"  # gives a Flag
        return out

    check(tmp_path, ["01-clean.pdf", "02-wrong-gst.pdf"], extract)
    ws = load_workbook(tmp_path / "review.xlsx").active
    headers = [c.value for c in ws[1]]
    cells = {h: [row[i] for row in ws.iter_rows(min_row=2)] for i, h in enumerate(headers)}

    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:O3"
    for i, header in enumerate(headers, start=1):
        width = ws.column_dimensions[get_column_letter(i)].width
        assert width >= len(header), header
    for header in ("Flags", "Notes", "Line items"):
        assert ws.column_dimensions[get_column_letter(headers.index(header) + 1)].width >= 40
        assert all(c.alignment.wrap_text for c in cells[header]), header
    needs_you, looks_fine = cells["Status"]  # Needs you rows first
    assert (needs_you.value, looks_fine.value) == ("Needs you", "Looks fine")
    assert looks_fine.fill.fgColor.rgb.endswith(GREEN)
    assert needs_you.fill.fgColor.rgb.endswith(RED)


def test_extract_returning_wrong_structure_gives_one_could_not_read_flag(tmp_path):
    def extract(pdf):
        out = fake_extract(pdf)
        del out["total"]
        out["supplier_name"] = "Southbridge Electrical Limited"  # bare value, no source text
        return out

    [row] = check(tmp_path, ["01-clean.pdf"], extract)
    assert row["Status"] == "Needs you"
    assert row["Flags"].startswith("Could not read:")
    assert "\n" not in row["Flags"]
    assert row["Supplier"] is None


@pytest.mark.parametrize("module, name", [(check_invoices.checks, "check"), (check_invoices.property_match, "suggest")])
def test_a_crash_in_the_checks_gives_that_invoice_one_could_not_read_flag(tmp_path, monkeypatch, module, name):
    real = getattr(module, name)

    def crash_on_invoice_2(values, *rest):
        if values["invoice_number"] == "INV-0582":
            raise KeyError("boom")
        return real(values, *rest)

    monkeypatch.setattr(module, name, crash_on_invoice_2)
    crashed, clean = check(tmp_path, ["01-clean.pdf", "02-wrong-gst.pdf"])  # Needs you rows first
    assert clean["Status"] == "Looks fine"
    assert crashed["Status"] == "Needs you"
    assert crashed["Flags"] == "Could not read: the checks failed on this invoice (KeyError: 'boom')."
    assert crashed["Suggested property"] is None


def test_low_confidence_is_a_note_and_status_stays_looks_fine(tmp_path):
    def extract(pdf):
        out = fake_extract(pdf)
        out["site"]["confidence"] = "low"
        return out

    [row] = check(tmp_path, ["01-clean.pdf"], extract)
    assert row["Status"] == "Looks fine"
    assert row["Flags"] is None
    assert row["Notes"] == "Low confidence: the reader was unsure about site 'Kowhai House'."


@pytest.mark.parametrize("text, amount", [
    ("$1,179.90", Decimal("1179.90")),
    ("1,850.00", Decimal("1850.00")),
    ("-$150.00", Decimal("-150.00")),
    ("-$22.50", Decimal("-22.50")),
])
def test_parse_amount_is_exact(text, amount):
    assert parse_amount(text) == amount
    assert isinstance(parse_amount(text), Decimal)


@pytest.mark.parametrize("text", ["", "about $100", "$1.2.3", "1179.905"])
def test_parse_amount_rejects_what_is_not_an_amount(text):
    with pytest.raises(ValueError):
        parse_amount(text)


@pytest.mark.parametrize("text, day", [
    ("03/04/2026", date(2026, 4, 3)),
    ("06/10/2026", date(2026, 10, 6)),
    ("9 October 2026", date(2026, 10, 9)),
    ("7 Oct 2026", date(2026, 10, 7)),
])
def test_parse_date_is_day_first(text, day):
    assert parse_date(text) == day


@pytest.mark.parametrize("text", ["31/02/2026", "13/13/2026", "next Tuesday", ""])
def test_parse_date_rejects_impossible_dates(text):
    with pytest.raises(ValueError):
        parse_date(text)
