"""Check contractor invoice PDFs and write review.xlsx, one row per invoice.

Usage: python check_invoices.py invoices/ [--open]
"""
import argparse
import os
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, ConfigDict, ValidationError
from pypdf import PdfReader

import checks
import property_match


class Strict(BaseModel):
    model_config = ConfigDict(extra="forbid")


class Source(Strict):
    source_text: str | None  # exactly as printed on the invoice, None if not there
    confidence: Literal["high", "medium", "low"]


class LineItem(Strict):
    description: Source
    amount: Source


class Extraction(Strict):
    """What the model must return. Values are worked out from source_text by code, never by the model."""
    document_type: Source
    supplier_name: Source
    gst_number: Source
    invoice_number: Source
    invoice_date: Source
    due_date: Source
    buyer_name: Source
    buyer_identifier: Source
    reference: Source
    work_order: Source
    site: Source
    line_items: list[LineItem]
    subtotal: Source
    gst: Source
    total: Source
    bank_account: Source


COLUMNS = [
    "File", "Status", "Flags", "Notes", "Suggested property", "Property reason",
    # Xero bill field order [assumed], with subtotal and GST next to total.
    "Supplier", "Date", "Due date", "Bill reference", "Line items", "Site", "Subtotal", "GST", "Total",
]
VALUE_COLUMNS = {
    "Supplier": "supplier_name", "Date": "invoice_date", "Due date": "due_date", "Bill reference": "invoice_number",
    "Site": "site", "Subtotal": "subtotal", "GST": "gst", "Total": "total",
}


def squash(text):
    """Ignore differences in spaces and line breaks, keep case."""
    return " ".join(text.split())


def parse_amount(text):
    """'$1,179.90' -> Decimal('1179.90'). Negative as printed on a credit note: '-$150.00'."""
    cleaned = re.sub(r"[\s$,]", "", text)
    if not re.fullmatch(r"-?\d+(\.\d{1,2})?", cleaned):
        raise ValueError(f"'{text}' is not an amount")
    return Decimal(cleaned)


def parse_date(text):
    """Day first: '03/04/2026' is 3 April 2026. Also '9 October 2026' and '7 Oct 2026'."""
    for fmt in ("%d/%m/%Y", "%d %B %Y", "%d %b %Y"):
        try:
            return datetime.strptime(squash(text), fmt).date()
        except ValueError:
            pass
    raise ValueError(f"'{text}' is not a valid date")


PARSERS = {"invoice_date": parse_date, "due_date": parse_date, "subtotal": parse_amount, "gst": parse_amount,
           "total": parse_amount}


def pdf_text(pdf):
    return "\n".join(page.extract_text() or "" for page in PdfReader(pdf).pages)


def check_invoice(pdf, extract):
    """One row for one PDF: status, Flags, Notes and the values worked out from verified source text."""
    flags, notes = [], []
    row = {"File": pdf.name, "Flags": flags, "Notes": notes}

    # Each failure below is one Flag and the rest is skipped, so every PDF still gets a row.
    try:
        text = squash(pdf_text(pdf))
    except Exception as e:
        flags.append(f"Unreadable: the PDF could not be opened ({e}).")
        return row
    if not text:
        flags.append("Unreadable: no text in the PDF, it may be a scan.")
        return row
    try:
        raw = extract(pdf)
    except Exception as e:
        error = re.sub(r"sk-[\w*-]+", "[key hidden]", str(e))  # OpenAI auth errors echo part of the key
        flags.append(f"Could not read: the invoice reader failed ({type(e).__name__}: {error}).")
        return row
    try:
        extraction = Extraction.model_validate(raw)
    except ValidationError as e:
        first = e.errors()[0]
        where = ".".join(str(part) for part in first["loc"])
        flags.append(f"Could not read: the invoice reader returned the wrong structure "
                     f"({e.error_count()} problems, first: {where}: {first['msg']}).")
        return row

    def read(name, found, parse=squash):
        label = name.replace("_", " ")
        if found.confidence == "low":
            shown = f"'{found.source_text}'" if found.source_text else "(nothing found)"
            notes.append(f"Low confidence: the reader was unsure about {label} {shown}.")
        if not squash(found.source_text or ""):  # only spaces or line breaks is the same as not there
            return None
        if squash(found.source_text) not in text:
            flags.append(f"Source text not found: {label} '{found.source_text}' is not in the PDF text.")
            return None
        try:
            return parse(found.source_text)
        except ValueError as e:
            flags.append(f"Could not read: {label} {e}.")
            return None

    values = {name: read(name, getattr(extraction, name), PARSERS.get(name, squash))
              for name in Extraction.model_fields if name != "line_items"}
    items = [(read("line item", i.description), read("line item amount", i.amount, parse_amount))
             for i in extraction.line_items]
    row.update({column: values[name] for column, name in VALUE_COLUMNS.items()})
    row["Line items"] = " | ".join(f"{d or ''} = {'' if a is None else f'${a:,.2f}'}" for d, a in items)
    row["values"], row["items"] = values, items  # for the checks, which also look across the run
    return row


WIDTHS = {
    "File": 26, "Status": 12, "Flags": 60, "Notes": 45, "Suggested property": 20, "Property reason": 30,
    "Supplier": 32, "Date": 12, "Due date": 12, "Bill reference": 15, "Line items": 60, "Site": 20,
    "Subtotal": 12, "GST": 12, "Total": 12,
}
WRAPPED = {"Flags", "Notes", "Line items"}
RED, GREEN = "FFC7CE", "C6EFCE"  # Excel's own Bad and Good fills
STATUS_STYLE = {
    "Needs you": (PatternFill("solid", fgColor=RED), Font(color="9C0006", bold=True)),
    "Looks fine": (PatternFill("solid", fgColor=GREEN), Font(color="006100", bold=True)),
}


def write_review(rows, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws.append(COLUMNS)
    for row in rows:
        ws.append([row.get(c) for c in COLUMNS])
    for header, column in zip(COLUMNS, ws.iter_cols()):
        ws.column_dimensions[column[0].column_letter].width = WIDTHS[header]
        column[0].font = Font(bold=True)
        for cell in column[1:]:
            cell.alignment = Alignment(vertical="top", wrap_text=header in WRAPPED)
            if isinstance(cell.value, Decimal):
                cell.number_format = "#,##0.00"
            elif isinstance(cell.value, date):
                cell.number_format = "DD/MM/YYYY"
            if header == "Status":
                cell.fill, cell.font = STATUS_STYLE[cell.value]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(output_path)


def run(folder, output_path, extract, run_date):
    """Check every PDF in folder and write one row per invoice to output_path."""
    rows = [check_invoice(pdf, extract) for pdf in sorted(Path(folder).glob("*.pdf"))]
    read = [r for r in rows if "values" in r]  # the rest already have their one Flag
    for row in read:
        others = [(r["File"], r["values"]) for r in read if r is not row]
        try:
            flags, notes = checks.check(row["values"], row["items"], run_date, others)
            found = (None, None, [])
            if not any(f.startswith("Not a bill:") for f in flags):  # not processed further
                found = property_match.suggest(row["values"], row["items"])
        except Exception as e:  # one bad invoice must not cost the others their rows
            row["Flags"].append(f"Could not read: the checks failed on this invoice ({type(e).__name__}: {e}).")
            continue
        row["Suggested property"], row["Property reason"], property_flags = found
        row["Flags"] += flags + property_flags
        row["Notes"] += notes
    for row in rows:
        row["Status"] = "Needs you" if row["Flags"] else "Looks fine"
        # One line each: a line break inside an error message must not look like a second Flag.
        row["Flags"] = "\n".join(squash(f) for f in row["Flags"]) or None
        row["Notes"] = "\n".join(squash(n) for n in row["Notes"]) or None
    rows.sort(key=lambda r: r["Status"] == "Looks fine")  # Needs you first; stable, so file order within each
    write_review(rows, output_path)
    return rows


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Check invoice PDFs and write review.xlsx.")
    parser.add_argument("folder", help="folder of invoice PDFs, e.g. invoices/")
    parser.add_argument("--open", action="store_true", help="open review.xlsx in Excel when the run finishes")
    args = parser.parse_args()
    from extractor import extract
    output = Path("review.xlsx").resolve()
    rows = run(args.folder, output, extract, date.today())
    needs_you = sum(r["Status"] == "Needs you" for r in rows)
    print(f"Wrote review.xlsx: {len(rows)} invoices, {needs_you} Needs you, {len(rows) - needs_you} Looks fine.")
    if args.open:
        os.startfile(output)  # Windows only: opens in the default app for .xlsx
